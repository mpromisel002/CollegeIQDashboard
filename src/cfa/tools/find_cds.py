"""Find Common Data Set documents for colleges that are not on the registry yet (owner tool).

    python -m cfa.tools.find_cds list --states PA,NY [--out todo.json]
    python -m cfa.tools.find_cds check --input candidates.json --out found.yaml [--report report.json]

`list` prints the four-year colleges in some states (IPEDS directory) that the registry lacks, with
each college's own web domain. Someone (a person, or an agent using a web search API) then adds
`urls`: pages or documents on that college's own site that look like its CDS.

`check` opens those pages (robots.txt respected, the college's own domain only), collects PDF/XLSX
links that look like a CDS, works out each one's year, and tries the newest first. A document is
kept only if it passes the registry rule (`validate_registry.passes`), so nothing unparseable is
ever proposed. Output is YAML ready to paste under `schools:` in registry/schools.yaml.

Search engines are never scraped (their terms forbid it); this tool only reads the colleges' sites.
"""

from __future__ import annotations

import argparse
import asyncio
import html
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlsplit, urlunsplit
from urllib.robotparser import RobotFileParser

from cfa import config
from cfa.sources.cds.adapter import parse_document
from cfa.sources.cds.fetch import fetch_cds
from cfa.sources.cds.parse_pdf import PDFIUM_LOCK
from cfa.sources.http import HttpClient
from cfa.sources.registry import load_bundled
from cfa.sources.urban_ipeds import is_candidate_college
from cfa.states import STATE_FIPS, parse_states
from cfa.tools.validate_registry import passes

OLDEST_YEAR = 2023  # a CDS older than 2023-2024 is too stale to list
# File hosts the registry does not accept: not the college's own site, and not direct downloads.
OFF_SITE = re.compile(
    r"drive\.google\.com|docs\.google\.com|sharepoint\.com|box\.com|dropbox\.com", re.IGNORECASE
)
CDS_TEXT = re.compile(r"common[\s_\-%20]*data[\s_\-%20]*set|(?<![a-z])cds(?![a-z])", re.IGNORECASE)
LINK = re.compile(r"<a\b[^>]*?href\s*=\s*[\"']([^\"'#]+)[\"'][^>]*>(.*?)</a>", re.IGNORECASE | re.DOTALL)
DOC_EXT = re.compile(r"\.(pdf|xlsx)(?:$|\?)", re.IGNORECASE)
# Year spellings seen in CDS file names and link text: 2025-2026, 2025_26, 2025-26, 25-26, 2526.
YEAR_PATTERNS = [
    re.compile(r"(?<!\d)(20\d\d)\s*[-_–/ ]\s*(20\d\d)(?!\d)"),
    re.compile(r"(?<!\d)(20\d\d)\s*[-_–/ ]\s*(\d\d)(?!\d)"),
    re.compile(r"(?<!\d)(\d\d)\s*[-_–]\s*(\d\d)(?!\d)"),
    re.compile(r"(?<![\d])(\d\d)(\d\d)(?![\d])"),
]
MAX_TRIES = 4  # documents tried per college, newest first
MAX_PAGES = 6  # HTML pages opened per college


def infer_year(text: str) -> str | None:
    """'CDS_2025-26.pdf' -> '2025-2026'. Only consecutive academic years count."""
    for i, pat in enumerate(YEAR_PATTERNS):
        for m in pat.finditer(text):
            a, b = m.group(1), m.group(2)
            start = int(a) if len(a) == 4 else 2000 + int(a)
            end = int(b) if len(b) == 4 else (start // 100) * 100 + int(b)
            if end == start + 1 and 2000 <= start <= 2040:
                if i == 3 and not CDS_TEXT.search(text):
                    continue  # a bare 4-digit "2526" only counts next to "CDS"
                return f"{start}-{end}"
    return None


def same_site(url: str, domain: str) -> bool:
    host = urlsplit(url).netloc.lower().split(":")[0]
    root = domain.lower().removeprefix("www.")
    return host == root or host.endswith("." + root)


def site_root(domain: str) -> str:
    """'www.cs.example.edu' -> 'example.edu': CDS files often live on another subdomain (ir.…, oira.…)."""
    parts = domain.lower().removeprefix("www.").split(".")
    return ".".join(parts[-2:]) if len(parts) >= 2 else domain


def https(url: str) -> str:
    sp = urlsplit(url)
    return urlunsplit(("https", sp.netloc, sp.path, sp.query, "")) if sp.scheme == "http" else url


@dataclass
class Candidate:
    url: str
    label: str = ""
    year: str | None = None

    @property
    def fmt(self) -> str | None:
        m = DOC_EXT.search(urlsplit(self.url).path + ("?" if urlsplit(self.url).query else ""))
        return m.group(1).lower() if m else None


@dataclass
class Outcome:
    unitid: int
    name: str
    entry: dict[str, Any] | None = None
    tried: list[str] = field(default_factory=list)
    reason: str = ""


def doc_links(page_url: str, text: str, domain: str) -> list[Candidate]:
    """PDF/XLSX links on a page that look like a CDS (by URL, link text, or a CDS page's year links)."""
    page_is_cds = bool(CDS_TEXT.search(page_url))
    out: list[Candidate] = []
    for href, raw_label in LINK.findall(text):
        label = html.unescape(" ".join(re.sub(r"<[^>]+>", " ", raw_label).split()))
        url = https(urljoin(page_url, html.unescape(href.strip())))
        if not url.startswith("https://") or not same_site(url, domain) or not DOC_EXT.search(url):
            continue
        hint = f"{url} {label}"
        if CDS_TEXT.search(hint) or (page_is_cds and infer_year(hint)):
            out.append(Candidate(url, label, infer_year(url) or infer_year(label)))
    return out


def year_pages(page_url: str, text: str, domain: str) -> list[str]:
    """Links to per-year CDS pages ("cds_2425.html", ".../cds-2025-2026/"), newest first."""
    found: list[tuple[str, str]] = []
    for href, raw_label in LINK.findall(text):
        label = html.unescape(" ".join(re.sub(r"<[^>]+>", " ", raw_label).split()))
        url = https(urljoin(page_url, html.unescape(href.strip())))
        if not url.startswith("https://") or not same_site(url, domain) or DOC_EXT.search(url):
            continue
        hint = f"{url} {label}"
        year = infer_year(hint)
        if CDS_TEXT.search(hint) and year and int(year[:4]) >= OLDEST_YEAR and url != page_url:
            found.append((year, url))
    return [u for _y, u in sorted(set(found), reverse=True)]


def rank(cands: list[Candidate]) -> list[Candidate]:
    """Newest year first; unknown years after dated ones; one entry per URL."""
    seen: dict[str, Candidate] = {}
    for c in cands:
        if c.url not in seen or (c.year and not seen[c.url].year):
            seen[c.url] = c
    usable = [c for c in seen.values() if c.year is None or int(c.year[:4]) >= OLDEST_YEAR]
    return sorted(usable, key=lambda c: (c.year is None, -(int(c.year[:4]) if c.year else 0)))


class Robots:
    def __init__(self, http: HttpClient) -> None:
        self.http = http
        self._cache: dict[str, RobotFileParser | None] = {}

    async def allowed(self, url: str) -> bool:
        host = urlsplit(url).netloc
        if host not in self._cache:
            res = await self.http.get(f"https://{host}/robots.txt")
            parser: RobotFileParser | None = None  # no robots.txt: allowed
            if res.ok and res.response is not None and res.response.status_code == 200:
                parser = RobotFileParser()
                parser.parse(res.response.text.splitlines())
            self._cache[host] = parser
        rules = self._cache[host]
        return rules is None or rules.can_fetch(config.USER_AGENT, url)


def document_year(content: bytes, fmt: str) -> str | None:
    """Read the year off the document's first page/sheet ("Common Data Set 2025-2026")."""
    try:
        if fmt == "pdf":
            import pypdfium2 as pdfium

            with PDFIUM_LOCK:  # pdfium is not thread-safe
                pdf = pdfium.PdfDocument(content)
                try:
                    text = " ".join(pdf[i].get_textpage().get_text_range() for i in range(min(2, len(pdf))))
                finally:
                    pdf.close()
        else:
            import io

            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                ws = wb.worksheets[0]
                text = " ".join(
                    str(c) for row in ws.iter_rows(max_row=40, values_only=True) for c in row if c
                )
            finally:
                wb.close()
    except Exception:
        return None
    m = re.search(r"common\s+data\s+set[^0-9]{0,40}(20\d\d\s*[-–/]\s*(?:20)?\d\d)", text, re.IGNORECASE)
    return infer_year(m.group(1)) if m else None


async def check_school(http: HttpClient, robots: Robots, item: dict[str, Any]) -> Outcome:
    unitid, name = int(item["unitid"]), str(item["name"])
    domain = site_root(str(item["domain"]))
    out = Outcome(unitid, name)
    cands: list[Candidate] = []
    pages = 0
    queue = [https(str(u)) for u in item.get("urls", [])]
    blocked = 0
    off_site = False
    while queue:
        url = queue.pop(0)
        if not same_site(url, domain):
            continue  # only the college's own site
        if DOC_EXT.search(url):
            cands.append(Candidate(url, "", infer_year(url)))
            continue
        if pages >= MAX_PAGES or not await robots.allowed(url):
            continue
        pages += 1
        res = await http.get(url)
        if res.response is not None and res.response.status_code in (401, 403):
            blocked += 1  # the site refuses automated requests; respected, never worked around
            continue
        ctype = res.response.headers.get("content-type", "") if res.response is not None else ""
        if res.ok and res.response is not None and "html" in ctype:
            final = str(res.response.url)
            cands += doc_links(final, res.response.text, domain)
            if CDS_TEXT.search(final) and OFF_SITE.search(res.response.text):
                off_site = True
            queue += [u for u in year_pages(final, res.response.text, domain)[:2] if u not in queue]
    ranked = rank(cands)
    if not ranked:
        if blocked:
            out.reason = "site refuses automated requests (HTTP 403)"
        elif off_site:
            out.reason = "CDS files are hosted off the college's site (Google Drive, SharePoint, Box…)"
        elif cands:
            out.reason = "newest CDS found is older than 2023-2024"
        else:
            out.reason = "no CDS document found"
        return out
    loop = asyncio.get_running_loop()
    for c in ranked[:MAX_TRIES]:
        out.tried.append(c.url)
        fmts = [c.fmt] if c.fmt else ["pdf", "xlsx"]
        content, fmt = None, ""
        for f in fmts:
            _status, content, note = await fetch_cds(http, c.url, f)
            if content is not None:
                fmt = f
                break
        if content is None:
            blocked_doc = "HTTP 403" in note or "HTTP 401" in note
            out.reason = (
                "site refuses automated requests (HTTP 403)" if blocked_doc else f"download failed: {note}"
            )
            continue
        year = c.year or document_year(content, fmt)
        if year is None:
            out.reason = "could not tell which year the document covers"
            continue
        if int(year[:4]) < OLDEST_YEAR:
            out.reason = f"newest document found is {year}"
            continue
        sections, err = await loop.run_in_executor(None, parse_document, content, fmt, year)
        if err:
            out.reason = f"unreadable: {err}"
            continue
        ok, why = passes(sections)
        if ok:
            out.entry = {"unitid": unitid, "name": name, "cds": [{"year": year, "url": c.url, "format": fmt}]}
            out.reason = ""
            return out
        out.reason = why
    return out


def drop_shared_documents(outcomes: list[Outcome]) -> list[Outcome]:
    """One document for two colleges (a system page, a main campus and its online arm) is ambiguous:
    neither gets it automatically. Documents already on the registry are not proposed again."""
    taken = {d.url for s in load_bundled().schools for d in s.cds}
    users: dict[str, list[Outcome]] = {}
    for o in outcomes:
        if o.entry:
            users.setdefault(o.entry["cds"][0]["url"], []).append(o)
    for url, group in users.items():
        if len(group) > 1 or url in taken:
            others = ", ".join(g.name for g in group)
            for o in group:
                o.entry = None
                o.reason = (
                    f"same document as {others if len(group) > 1 else 'a registry school'}; choose by hand"
                )
    return outcomes


def entries_yaml(entries: list[dict[str, Any]]) -> str:
    lines = []
    for e in entries:
        d = e["cds"][0]
        lines += [
            f"  - unitid: {e['unitid']}",
            f"    name: {json.dumps(e['name'])}",
            "    cds:",
            f'      - year: "{d["year"]}"',
            f"        url: {d['url']}",
            f"        format: {d['format']}",
        ]
    return "\n".join(lines) + ("\n" if lines else "")


async def list_colleges(states: list[str]) -> list[dict[str, Any]]:
    listed = {s.unitid for s in load_bundled().schools}
    codes = [STATE_FIPS[s] for s in states] or sorted(STATE_FIPS.values())
    out: list[dict[str, Any]] = []
    async with HttpClient() as http:
        for i in range(0, len(codes), 10):
            fips = ",".join(str(c) for c in codes[i : i + 10])
            for year in range(config.current_year() - 1, config.current_year() - 4, -1):
                res = await http.get(f"{config.URBAN_BASE_URL}/directory/{year}/?fips={fips}")
                rows = res.json().get("results", []) if res.ok else []
                if rows:
                    break
            for r in rows:
                if not is_candidate_college(r) or r.get("unitid") in listed:
                    continue
                site = str(r.get("url_school") or "")
                domain = urlsplit(site if "//" in site else f"//{site}").netloc.lower().removeprefix("www.")
                if domain:
                    out.append(
                        {
                            "unitid": r["unitid"],
                            "name": r["inst_name"],
                            "state": r.get("state_abbr"),
                            "domain": domain,
                        }
                    )
    return sorted(out, key=lambda r: (str(r["state"]), str(r["name"])))


async def check_all(items: list[dict[str, Any]], workers: int = 6) -> list[Outcome]:
    async with HttpClient(max_retries=1) as http:
        robots = Robots(http)
        sem = asyncio.Semaphore(workers)
        done = 0

        async def one(item: dict[str, Any]) -> Outcome:
            nonlocal done
            async with sem:
                try:
                    res = await check_school(http, robots, item)
                except Exception as exc:  # one odd site must not stop the batch
                    res = Outcome(
                        int(item["unitid"]), str(item["name"]), reason=f"error: {type(exc).__name__}"
                    )
            done += 1
            mark = "FOUND" if res.entry else "-"
            print(f"[{done}/{len(items)}] {mark:5} {res.name[:50]} {res.reason}", file=sys.stderr, flush=True)
            return res

        return list(await asyncio.gather(*(one(i) for i in items)))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    lp = sub.add_parser("list", help="colleges in some states that the registry lacks")
    lp.add_argument("--states", default="", help="e.g. PA,NY (default: every state)")
    lp.add_argument("--out", type=Path)
    cp = sub.add_parser("check", help="validate candidate URLs and write registry entries")
    cp.add_argument("--input", type=Path, required=True)
    cp.add_argument("--out", type=Path, required=True, help="YAML entries for registry/schools.yaml")
    cp.add_argument("--report", type=Path, help="JSON with every college's outcome")
    args = ap.parse_args(argv)
    if args.cmd == "list":
        codes, unknown = parse_states(args.states)
        if unknown:
            ap.error(f"not states: {', '.join(unknown)}")
        rows = asyncio.run(list_colleges(codes))
        text = json.dumps(rows, indent=1)
        if args.out:
            args.out.write_text(text, encoding="utf-8")
            print(f"{len(rows)} colleges written to {args.out}")
        else:
            print(text)
        return 0
    items = json.loads(args.input.read_text(encoding="utf-8"))
    outcomes = drop_shared_documents(asyncio.run(check_all(items)))
    found = [o.entry for o in outcomes if o.entry]
    args.out.write_text(entries_yaml(found), encoding="utf-8")
    if args.report:
        args.report.write_text(json.dumps([o.__dict__ for o in outcomes], indent=1), encoding="utf-8")
    print(f"{len(found)}/{len(outcomes)} colleges have a CDS that passes; entries in {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
