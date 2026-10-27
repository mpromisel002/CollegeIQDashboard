"""CDS XLSX -> layout Document. Values are located by label text, not cell coordinates."""

from __future__ import annotations

import datetime as dt
import io
import re
from typing import Any

from cfa.sources.cds.layout import Document, Row, Token

COL_WIDTH = 10.0  # synthetic x units per column so the same proximity logic works
# The official template's machine-readable "answer sheet" columns start with codes like C.101 / C.8G04.
ANSWER_CODE = re.compile(r"^[A-J]\.\s?\d")
MAX_ROWS_PER_SHEET = 3000
# Sheets plainly titled as a part the extractors never read ("CDS-G", "I. Faculty", "Section J").
_UNUSED_PART = re.compile(r"^(?:cds[\s_-]*|section[\s_-]*)?([adgij])(?:[\s._:-]|$)", re.IGNORECASE)


class XlsxError(Exception):
    pass


def _fmt(value: Any, number_format: str | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dt.datetime | dt.date | dt.time):
        return None
    if isinstance(value, int | float):
        if number_format and "%" in number_format:
            return f"{round(value * 100, 4):g}%"
        return f"{value:g}" if isinstance(value, float) else str(value)
    s = str(value).strip()
    return s or None


def parse_xlsx(content: bytes) -> Document:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise XlsxError(f"Not a readable XLSX ({type(exc).__name__})") from exc
    rows: list[Row] = []
    try:
        for sheet_no, ws in enumerate(wb.worksheets):
            title = ws.title.strip()
            if title.lower().startswith(("answer", "definition", "cds definition", "glossary")):
                continue
            if _UNUSED_PART.match(title):
                continue
            for r_idx, row in enumerate(ws.iter_rows()):
                if r_idx >= MAX_ROWS_PER_SHEET:
                    break
                tokens: list[Token] = []
                for c_idx, cell in enumerate(row):
                    text = _fmt(getattr(cell, "value", None), getattr(cell, "number_format", None))
                    if text is None:
                        continue
                    if ANSWER_CODE.match(text):
                        break
                    x = c_idx * COL_WIDTH
                    words = text.split()
                    for w_i, w in enumerate(words):
                        # keep words of one cell in order, inside the cell's column span
                        tokens.append(Token(w, x + w_i * 0.01, x + COL_WIDTH * 0.9))
                if tokens:
                    rows.append(Row(tokens, float(r_idx), sheet_no))
    finally:
        wb.close()
    if len(rows) < 50:
        raise XlsxError("Workbook does not look like a Common Data Set")
    return Document(rows, "xlsx")
